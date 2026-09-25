#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
GES (Greedy Equivalence Search) desfasado t -> t+1 sobre 'FCM DATOS.xlsx'.

SOLO ESTRUCTURA: estos algoritmos responden si HAY o NO relación, no cuánta.
PC, FCI y GES no devuelven ni signo ni valor de la arista; el peso se obtiene
después, en una segunda fase, entrenando un modelo con la estructura ya fijada.
Por eso la celda marca la arista y nada más.

  - Descubrimiento causal desfasado (lag 1 día): para cada objetivo j se construye
    el conjunto [ X(t) de las 11 variables | j(t+1) ] y se corre GES (score BIC).
    Las variables adyacentes a j(t+1) en el CPDAG rellenan la columna j.
  - Celda i->j = '→'  si GES identifica orientacion para esa arista.
                 '--' si halla la relacion pero deja la arista sin orientar
                      (no identificable dentro de la clase de equivalencia).
                 vacio si no hay relacion.
  - El SENTIDO de la flecha lo impone el desfase temporal (t precede a t+1), no
    GES: el algoritmo desconoce el orden temporal y de hecho orienta la mayoria
    de estas aristas al reves. Lo que se conserva de GES es SI la orientacion es
    identificable o no, que es lo que distingue '→' de '--'.
  - Diagonal E->E = persistencia (la variable en t sobre si misma en t+1).
  - Objetivos: solo los niveles E1, E3, E5 y E7. Las tendencias E2, E4, E6 son
    derivadas deterministas y no se predicen.

Escribe dos hojas nuevas en 'FCM DATOS.xlsx', clonando el formato de las de PC:
  - PROXIMOS ESTADOS - GES
  - PREDECIR CONTAGIO - GES   (solo la columna E1)

Uso:
    python3 ges_datos_excel.py
"""
import os
import numpy as np
import pandas as pd
from causallearn.search.ScoreBased.GES import ges

from pc_fci import cargar, SHORT

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, "..", "datos", "FCM DATOS.xlsx")
TAB = os.path.join(HERE, "..", "resultados", "ges", "tablas")
os.makedirs(TAB, exist_ok=True)

PAISES = ["ES", "GB", "BR"]
OBJETIVOS = ["E1", "E3", "E5", "E7"]      # niveles; las tendencias no se predicen

# hojas plantilla (de las que se clona el formato) y hojas destino
PLANTILLA_PROX, HOJA_PROX = "PROXIMOS ESTADOS - PC", "PROXIMOS ESTADOS - GES"
PLANTILLA_PRED, HOJA_PRED = "PREDECIR CONTAGIO - PC", "PREDECIR CONTAGIO - GES"

# fila de la etiqueta de país en cada plantilla (el resto se deduce)
FILAS_PROX = {"ES": 3, "GB": 19, "BR": 36}
FILAS_PRED = {"ES": 3, "GB": 20, "BR": 38}

LEYENDA = ("Celda i→j:  «→» hay relación y GES identifica su orientación;  «--» SOBRE FONDO VERDE: hay relación pero "
           "el algoritmo no sabe en qué dirección;  CELDA VACÍA: no hay relación (que es distinto de no saber la "
           "dirección). SOLO ESTRUCTURA: PC, FCI y GES no devuelven "
           "signo ni valor de la arista —responden si hay o no relación, no cuánta—; el peso se obtiene después "
           "entrenando un modelo sobre esta estructura. El sentido de la flecha lo fija el desfase temporal (t precede "
           "a t+1), no el algoritmo. Diagonal E→E = persistencia. ")

NOTA_PROX = ("Descubrimiento causal desfasado t → t+1 (lag 1 día) con GES (Greedy Equivalence Search, score BIC): "
             "para cada objetivo j se corre GES sobre [X(t) | j(t+1)] y se toman las variables adyacentes a j(t+1). "
             + LEYENDA +
             "Objetivos: niveles E1,E3,E5 y E7 (las tendencias E2,E4,E6 son derivadas deterministas y no se predicen).")
NOTA_PRED = ("Descubrimiento causal desfasado t → t+1 con GES (score BIC). Solo la columna E1 (contagio): qué variables "
             "en t tienen relación con el contagio en t+1. " + LEYENDA)


# ----------------------------------------------------------------------
DIRIGIDA, SIN_DIR = "→", "--"


def ges_desfasado(cc):
    """Devuelve A[i][j] = '→' | '--' para las aristas i(t) — j(t+1) que halla GES."""
    X, labels, fuera, n = cargar(cc)
    L = len(labels)
    idx = {lab: k for k, lab in enumerate(labels)}
    Xt, Xt1 = X[:-1], X[1:]

    A = {i: {} for i in labels}
    sel = {}
    for j in OBJETIVOS:
        if j not in idx:
            continue
        D = np.hstack([Xt, Xt1[:, [idx[j]]]])          # 11 predictores en t + objetivo en t+1
        G = ges(D, score_func="local_score_BIC")["G"].graph
        marcas = {}
        for i in range(L):
            a, b = G[i, L], G[L, i]
            if a == 0 and b == 0:
                continue
            # (-1,1) y (1,-1) son aristas orientadas por GES; (-1,-1) queda sin orientar.
            # El sentido que da GES se descarta: lo fija el orden temporal t → t+1.
            marcas[labels[i]] = SIN_DIR if (a == -1 and b == -1) else DIRIGIDA
        sel[j] = marcas
        for i, m in marcas.items():
            A[i][j] = m
    return A, sel, labels, fuera, n


# ----------------------------------------------------------------------
def rellenar(ws, fila_pais, A, labels, solo_e1, nota):
    """Sobrescribe un bloque de país ya formateado (clonado de la hoja de PC)."""
    from openpyxl.styles import Font, PatternFill

    r0 = fila_pais + 2                                  # primera fila de datos
    AZUL, VERDE_TXT = "1F4E79", "1E6C2F"
    # verde = hay relación pero el algoritmo no identifica la dirección.
    # Se distingue así de la celda vacía, que significa que NO hay relación.
    VERDE = PatternFill("solid", fgColor="C6EFCE")
    for fi, i in enumerate(SHORT):
        for ci, j in enumerate(SHORT):
            c = ws.cell(r0 + fi, 2 + ci)
            c.value = None                              # limpia el valor de PC
            if solo_e1 and j != "E1":
                continue
            if j not in OBJETIVOS or i not in labels or j not in labels:
                continue
            m = A.get(i, {}).get(j)
            if m is None:
                continue
            c.value = m
            if m == DIRIGIDA:
                c.font = Font(bold=True, size=11, color=AZUL)
            else:
                c.font = Font(bold=True, size=11, color=VERDE_TXT)
                c.fill = VERDE
    ws.cell(fila_pais + 14, 1).value = nota             # fila de nota del bloque


def main():
    import openpyxl

    res = {cc: ges_desfasado(cc) for cc in PAISES}
    for cc in PAISES:
        A, sel, labels, fuera, n = res[cc]
        print(f"{cc}: n={n}  vars={len(labels)}"
              + (f"  [excluidas: {', '.join(fuera)}]" if fuera else ""))
        for j in OBJETIVOS:
            m = sel.get(j, {})
            print(f"    {j}(t+1): " + ("  ".join(f"{i}{m[i]}" for i in sorted(m)) or "(nada)"))

    wb = openpyxl.load_workbook(XLSX)
    for plantilla, hoja, filas, solo_e1, nota, titulo, desc in [
        (PLANTILLA_PROX, HOJA_PROX, FILAS_PROX, False, NOTA_PROX,
         "PROXIMOS ESTADOS (t→t+1) — GES",
         "Matriz de próximos estados descubierta por el algoritmo GES, por país (ES, GB, BR)."),
        (PLANTILLA_PRED, HOJA_PRED, FILAS_PRED, True, NOTA_PRED,
         "PREDECIR CONTAGIO (t→t+1) — GES",
         "Predictores del contagio E1 en t+1 (columna E1) descubiertos por GES, por país."),
    ]:
        if hoja in wb.sheetnames:
            del wb[hoja]
        ws = wb.copy_worksheet(wb[plantilla])           # clona formato, anchos y merges
        ws.title = hoja
        ws["A1"] = titulo
        ws["A2"] = desc
        for cc, fila in filas.items():
            A, sel, labels, fuera, n = res[cc]
            rellenar(ws, fila, A, labels, solo_e1, nota)

    wb.save(XLSX)

    filas_csv = []
    for cc in PAISES:
        A, sel, labels, fuera, n = res[cc]
        for j in OBJETIVOS:
            for i, m in sorted(sel.get(j, {}).items()):
                filas_csv.append({"pais": cc, "desde": i, "hacia": j,
                                  "arista": "dirigida" if m == DIRIGIDA else "sin_direccion",
                                  "relacion": "persistencia" if i == j else "cruzada"})
    pd.DataFrame(filas_csv).to_csv(os.path.join(TAB, "GES_desfasado.csv"), index=False)

    print(f"\nHojas escritas en {os.path.basename(XLSX)}: '{HOJA_PROX}' y '{HOJA_PRED}'")
    print("Tabla: resultados/ges/tablas/GES_desfasado.csv")


if __name__ == "__main__":
    main()
